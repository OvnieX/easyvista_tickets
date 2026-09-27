#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
===============================================================================
 EasyVista - Tickets uniques + Dashboard Excel
===============================================================================

OBJECTIF
--------
Le fichier "Mes actions" d'EasyVista contient une ligne par ACTION.
Par conséquent, un même ticket peut apparaître plusieurs fois.

Ce script :
  1. lit un export EasyVista au format CSV ou XLSX ;
  2. nettoie les particularités de l'export EasyVista ;
  3. supprime les doublons pour ne conserver qu'UNE ligne par ticket ;
  4. génère un nouveau fichier Excel contenant :
       - un onglet "Tickets uniques" ;
       - un onglet "Dashboard" avec plusieurs graphiques.

IMPORTANT
---------
Le script N'AJOUTE PAS de colonne "Nombre d'actions".

Lorsqu'un ticket apparaît plusieurs fois, le script choisit la ligne la plus
récente avec les règles suivantes :
  - d'abord la "Date de dernière modification" ;
  - en cas d'égalité, la "Date de fin réelle" ;
  - puis la "Date d'émission".

Cela évite de conserver arbitrairement la première action d'un ticket lorsque
plusieurs lignes ont exactement la même date de dernière modification.

DEPENDANCES
-----------
Pour un fichier CSV :
    pip install xlsxwriter

Pour lire également les fichiers XLSX :
    pip install xlsxwriter openpyxl

EXEMPLES
--------
Traitement d'un CSV EasyVista :
    python easyvista_tickets.py "Mes actions.csv"

Traitement d'un fichier Excel :
    python easyvista_tickets.py "Mes actions.xlsx"

Choisir le fichier de sortie :
    python easyvista_tickets.py "Mes actions.csv" -o "Tickets_EasyVista.xlsx"

Afficher les informations techniques de lecture :
    python easyvista_tickets.py "Mes actions.csv" --verbose

Pour un XLSX, choisir explicitement une feuille :
    python easyvista_tickets.py "Mes actions.xlsx" --sheet "Mes actions"
===============================================================================
"""

# =============================================================================
# 1) IMPORTS
# =============================================================================
# Les modules csv, argparse, pathlib, datetime... font partie de Python.
# xlsxwriter et openpyxl sont importés plus bas uniquement quand ils sont utiles.

import argparse
import csv
import re
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path


# =============================================================================
# 2) CONFIGURATION : NOMS DES COLONNES EASYVISTA
# =============================================================================
# Centraliser les noms ici permet d'adapter facilement le script si EasyVista
# renomme un jour une colonne dans l'export.

TICKET_COL = "N°"
DATE_MODIF_COL = "Date de dernière modification"
DATE_FIN_COL = "Date de fin réelle"
DATE_EMISSION_COL = "Date d'émission"

# Encodages testés pour les CSV. L'export actuel est en UTF-8 avec BOM, mais
# CP1252 est courant sur des exports Windows/Excel français.
CSV_ENCODINGS = ("utf-8-sig", "utf-8", "cp1252", "latin-1")


# =============================================================================
# 3) FONCTIONS UTILITAIRES
# =============================================================================

def parse_date(value):
    """Convertit une date EasyVista en objet datetime.

    Plusieurs formats sont acceptés afin de ne pas dépendre d'un format unique.
    Si la valeur est vide ou invalide, datetime.min est renvoyé. Cela permet de
    comparer les dates sans provoquer d'erreur.
    """

    value = str(value or "").strip()

    formats = (
        "%d/%m/%Y %H:%M:%S",  # ex. 21/08/2026 03:32:44
        "%d/%m/%Y %H:%M",     # ex. 21/08/2026 03:32
        "%d/%m/%Y",            # ex. 21/08/2026
    )

    for fmt in formats:
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue

    return datetime.min


def clean_value(value, default="Non renseigné"):
    """Nettoie une valeur texte pour le Dashboard."""

    text = str(value or "").strip()
    return text if text else default


def natural_key(value):
    """Clé de tri 'naturel'.

    Exemple : 2 sera classé avant 10, contrairement à un tri texte classique.
    Utile notamment pour les priorités.
    """

    parts = re.split(r"(\d+)", str(value))
    return [int(part) if part.isdigit() else part.lower() for part in parts]


def excel_col_name(index):
    """Convertit un index de colonne 0-based en lettre Excel.

    0 -> A, 1 -> B, 25 -> Z, 26 -> AA, etc.
    """

    index += 1
    letters = []

    while index:
        index, remainder = divmod(index - 1, 26)
        letters.append(chr(65 + remainder))

    return "".join(reversed(letters))


# =============================================================================
# 4) LECTURE ET NORMALISATION DE L'EXPORT EASYVISTA
# =============================================================================

def normalize_row(row, width, line_no=None):
    """Ramène une ligne au même nombre de colonnes que l'en-tête.

    Pourquoi cette fonction existe ?
    --------------------------------
    L'export EasyVista utilisé ici possède 22 colonnes dans l'en-tête, mais
    certaines lignes contiennent un 23e champ VIDE à la fin.

    C'est précisément ce qui provoquait l'erreur :
        Expected 22 fields in line 3, saw 23

    Comportement :
      - les champs supplémentaires VIDES situés à droite sont supprimés ;
      - si un champ supplémentaire NON VIDE existe, le script s'arrête plutôt
        que de supprimer silencieusement une information ;
      - une ligne trop courte est complétée avec des cellules vides.
    """

    row = list(row)

    # Supprime uniquement les colonnes supplémentaires vides à droite.
    while len(row) > width and str(row[-1]).strip() == "":
        row.pop()

    # S'il reste encore des colonnes en trop, elles contiennent des données.
    # On préfère donc signaler clairement le problème.
    if len(row) > width:
        extra = row[width:]
        raise ValueError(
            f"Ligne {line_no or '?'} : {len(row)} champs pour {width} colonnes. "
            f"Champs supplémentaires non vides : {extra!r}"
        )

    # Si une ligne possède moins de colonnes que l'en-tête, on complète.
    if len(row) < width:
        row.extend([""] * (width - len(row)))

    return row


def detect_csv_encoding(path):
    """Détermine un encodage lisible pour le CSV.

    On teste les encodages les plus probables. Le fichier n'est pas encore
    analysé ici : on vérifie seulement qu'il peut être décodé sans erreur.
    """

    raw = path.read_bytes()

    for encoding in CSV_ENCODINGS:
        try:
            raw.decode(encoding)
            return encoding
        except UnicodeDecodeError:
            continue

    # latin-1 peut techniquement décoder n'importe quel octet ; ce cas est donc
    # très improbable, mais le message est plus clair si la liste change.
    raise ValueError("Impossible de déterminer l'encodage du fichier CSV.")


def read_csv_easyvista(path, verbose=False):
    """Lit un export CSV EasyVista de façon tolérante."""

    encoding = detect_csv_encoding(path)

    with path.open("r", encoding=encoding, newline="") as file:
        # On prélève un échantillon pour détecter le séparateur.
        sample = file.read(8192)
        file.seek(0)

        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
        except csv.Error:
            # L'export EasyVista français utilise normalement le point-virgule.
            delimiter = ";"

        reader = csv.reader(file, delimiter=delimiter, quotechar='"')
        rows = list(reader)

    if not rows:
        raise ValueError("Le fichier CSV est vide.")

    # Première ligne = noms des colonnes.
    headers = [str(value).strip() for value in rows[0]]

    if not any(headers):
        raise ValueError("La ligne d'en-tête du CSV est vide.")

    width = len(headers)
    data = []

    # Toutes les lignes suivantes sont converties en dictionnaires :
    # {"N°": "I0963193", "Titre": "...", ...}
    for line_no, raw_row in enumerate(rows[1:], start=2):
        # Ignore les lignes entièrement vides.
        if not raw_row or not any(str(value).strip() for value in raw_row):
            continue

        row = normalize_row(raw_row, width, line_no)
        data.append(dict(zip(headers, row)))

    if verbose:
        printable_delimiter = "TAB" if delimiter == "\t" else delimiter
        print(f"[Lecture CSV] Encodage  : {encoding}")
        print(f"[Lecture CSV] Séparateur: {printable_delimiter!r}")
        print(f"[Lecture CSV] Colonnes  : {width}")
        print(f"[Lecture CSV] Lignes    : {len(data)}")

    return headers, data


def read_xlsx(path, sheet_name=None, verbose=False):
    """Lit un export XLSX EasyVista.

    openpyxl est utilisé uniquement pour LIRE le fichier source Excel.
    Le fichier de résultat est créé avec xlsxwriter.
    """

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "Module manquant pour lire un XLSX : pip install openpyxl"
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)

    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(
                    f"Feuille '{sheet_name}' introuvable. "
                    f"Feuilles disponibles : {', '.join(workbook.sheetnames)}"
                )
            worksheet = workbook[sheet_name]
        else:
            worksheet = workbook.active

        iterator = worksheet.iter_rows(values_only=True)

        try:
            headers = [str(value or "").strip() for value in next(iterator)]
        except StopIteration:
            raise ValueError("Le fichier Excel est vide.")

        if not any(headers):
            raise ValueError("La ligne d'en-tête du fichier Excel est vide.")

        width = len(headers)
        data = []

        for line_no, raw_row in enumerate(iterator, start=2):
            row = ["" if value is None else str(value) for value in raw_row]

            if not any(value.strip() for value in row):
                continue

            row = normalize_row(row, width, line_no)
            data.append(dict(zip(headers, row)))

        if verbose:
            print(f"[Lecture XLSX] Feuille  : {worksheet.title}")
            print(f"[Lecture XLSX] Colonnes : {width}")
            print(f"[Lecture XLSX] Lignes   : {len(data)}")

        return headers, data

    finally:
        # Le finally garantit la fermeture du fichier même en cas d'erreur.
        workbook.close()


def read_export(path, sheet_name=None, verbose=False):
    """Choisit automatiquement la bonne méthode selon l'extension."""

    extension = path.suffix.lower()

    if extension == ".csv":
        return read_csv_easyvista(path, verbose=verbose)

    if extension == ".xlsx":
        return read_xlsx(path, sheet_name=sheet_name, verbose=verbose)

    raise ValueError("Formats acceptés : .csv ou .xlsx")


# =============================================================================
# 5) FILTRAGE ET DEDOUBLONNAGE DES TICKETS
# =============================================================================

def valid_ticket(value):
    """Indique si la valeur ressemble à un vrai numéro de ticket.

    EasyVista ajoute dans l'export une ligne technique ressemblant à :
        - ; 2 ; 2 ; 2 ; ...

    Cette ligne doit être ignorée.
    """

    value = str(value or "").strip()

    if not value or value == "-":
        return False

    # Une valeur d'un seul caractère (ex. "2") est considérée comme technique.
    return len(value) > 1


def representative_sort_key(row):
    """Retourne la clé utilisée pour choisir la ligne d'un ticket en doublon.

    Ordre de priorité :
      1. Date de dernière modification ;
      2. Date de fin réelle ;
      3. Date d'émission.

    Le deuxième critère est important avec "Mes actions" : plusieurs actions
    d'un même ticket ont souvent exactement la même date de dernière
    modification. La date de fin réelle permet alors de garder l'action la plus
    récente au lieu de la première ligne rencontrée dans le CSV.
    """

    return (
        parse_date(row.get(DATE_MODIF_COL)),
        parse_date(row.get(DATE_FIN_COL)),
        parse_date(row.get(DATE_EMISSION_COL)),
    )


def deduplicate(headers, rows, verbose=False):
    """Retourne les lignes valides et une seule ligne par numéro de ticket."""

    if TICKET_COL not in headers:
        raise KeyError(
            f"Colonne '{TICKET_COL}' absente. "
            f"Colonnes trouvées : {', '.join(headers)}"
        )

    # Retire notamment la ligne technique EasyVista '-;2;2;2;...'.
    valid_rows = [row for row in rows if valid_ticket(row.get(TICKET_COL))]

    # Dictionnaire : numéro du ticket -> ligne représentative.
    unique = {}

    for row in valid_rows:
        ticket = str(row.get(TICKET_COL, "")).strip()
        previous = unique.get(ticket)

        # Première apparition du ticket : on la conserve provisoirement.
        if previous is None:
            unique[ticket] = row
            continue

        # Si une autre ligne du même ticket est plus récente, elle remplace
        # la ligne précédemment conservée.
        if representative_sort_key(row) > representative_sort_key(previous):
            unique[ticket] = row

    tickets = list(unique.values())

    # Pour l'utilisateur, la liste finale est plus pratique avec les tickets
    # récemment émis en premier.
    tickets.sort(
        key=lambda row: parse_date(row.get(DATE_EMISSION_COL)),
        reverse=True,
    )

    if verbose:
        duplicates = len(valid_rows) - len(tickets)
        print(f"[Dédoublonnage] Lignes valides : {len(valid_rows)}")
        print(f"[Dédoublonnage] Tickets uniques: {len(tickets)}")
        print(f"[Dédoublonnage] Doublons retirés: {duplicates}")

    return valid_rows, tickets


# =============================================================================
# 6) PREPARATION DES DONNEES DU DASHBOARD
# =============================================================================

def readable_type(row):
    """Détermine si le ticket est un Incident ou une Demande de service.

    Dans l'export actuel, la colonne "Type" contient parfois le nom d'une icône
    (ex. DB_Incident.png ou service.png). On utilise plusieurs indices afin
    d'obtenir un libellé lisible dans le graphique.
    """

    type_value = clean_value(row.get("Type"))
    type_lower = type_value.lower()

    # 1) Information directement présente dans la colonne Type.
    if "incident" in type_lower:
        return "Incident"

    if "service" in type_lower or "demande" in type_lower:
        return "Demande de service"

    # 2) Information souvent présente dans "Sujet (complet)".
    subject = str(row.get("Sujet (complet)") or "").strip().lower()

    if subject.startswith("incidents/"):
        return "Incident"

    if subject.startswith("demande de service/"):
        return "Demande de service"

    # 3) Dernier filet de sécurité adapté à la numérotation observée :
    # Ixxxxxxx = Incident ; Sxxxxxxx = Service Request.
    ticket = str(row.get(TICKET_COL) or "").strip().upper()

    if ticket.startswith("I"):
        return "Incident"

    if ticket.startswith("S"):
        return "Demande de service"

    # Si aucun indice ne permet de décider, on garde la valeur d'origine.
    return type_value


def build_dashboard_data(tickets):
    """Calcule les regroupements utilisés par les tableaux et graphiques."""

    # Nombre de TICKETS UNIQUES par statut / priorité / groupe / type.
    status = Counter(clean_value(row.get("Statut")) for row in tickets)
    priority = Counter(clean_value(row.get("Priorité")) for row in tickets)
    groups = Counter(clean_value(row.get("Groupe")) for row in tickets)
    types = Counter(readable_type(row) for row in tickets)

    # Regroupement mensuel selon la date d'émission du ticket.
    # La clé YYYY-MM est volontaire : elle se trie naturellement par date.
    months = Counter()

    for row in tickets:
        date_emission = parse_date(row.get(DATE_EMISSION_COL))

        if date_emission != datetime.min:
            months[date_emission.strftime("%Y-%m")] += 1

    return status, priority, types, groups, months


# =============================================================================
# 7) CREATION DU FICHIER EXCEL FINAL
# =============================================================================

def write_output(headers, tickets, output):
    """Crée le classeur Excel final avec la liste et le Dashboard."""

    try:
        import xlsxwriter
        from xlsxwriter.exceptions import FileCreateError
    except ImportError as exc:
        raise RuntimeError(
            "Module manquant pour créer le fichier Excel : pip install xlsxwriter"
        ) from exc

    # Le dossier de sortie est créé automatiquement s'il n'existe pas.
    output.parent.mkdir(parents=True, exist_ok=True)

    try:
        workbook = xlsxwriter.Workbook(str(output))

        # Deux feuilles : données dédoublonnées + synthèse graphique.
        ws = workbook.add_worksheet("Tickets uniques")
        dash = workbook.add_worksheet("Dashboard")

        # ---------------------------------------------------------------------
        # 7.1) Formats Excel réutilisables
        # ---------------------------------------------------------------------
        fmt_header = workbook.add_format({
            "bold": True,
            "font_color": "white",
            "bg_color": "#1F4E78",
            "align": "center",
            "valign": "vcenter",
            "border": 1,
        })

        fmt_title = workbook.add_format({
            "bold": True,
            "font_color": "white",
            "bg_color": "#1F4E78",
            "font_size": 16,
            "align": "center",
            "valign": "vcenter",
        })

        fmt_section = workbook.add_format({
            "bold": True,
            "bg_color": "#D9EAF7",
            "border": 1,
        })

        fmt_kpi = workbook.add_format({
            "bold": True,
            "font_size": 14,
            "align": "center",
        })

        fmt_wrap = workbook.add_format({
            "text_wrap": True,
            "valign": "top",
        })

        fmt_center = workbook.add_format({
            "align": "center",
            "valign": "top",
        })

        # ---------------------------------------------------------------------
        # 7.2) Onglet "Tickets uniques"
        # ---------------------------------------------------------------------

        # Ligne d'en-tête.
        for column_index, header in enumerate(headers):
            ws.write(0, column_index, header, fmt_header)

        # Données : une ligne par ticket.
        # Certaines colonnes courtes sont centrées, les autres restent en texte
        # avec retour à la ligne pour conserver une bonne lisibilité.
        centered_columns = {
            "N°", "Statut", "Priorité", "Impact", "Urgence", "Type"
        }

        for row_index, row in enumerate(tickets, start=1):
            for column_index, header in enumerate(headers):
                value = row.get(header, "")
                cell_format = fmt_center if header in centered_columns else fmt_wrap
                ws.write(row_index, column_index, value, cell_format)

        # Fige la ligne des titres lorsqu'on fait défiler le fichier.
        ws.freeze_panes(1, 0)

        # Active le filtre Excel sur toutes les colonnes.
        if headers:
            ws.autofilter(0, 0, max(len(tickets), 1), len(headers) - 1)

        # Largeur spécifique des colonnes importantes.
        widths = {
            "N°": 14,
            "Titre": 52,
            "Statut": 22,
            "Priorité": 12,
            "Impact": 16,
            "Urgence": 16,
            "Groupe": 24,
            "Demandeur": 24,
            "Type": 20,
            "Sujet": 28,
            "Sujet (complet)": 44,
            "Bénéficiaire": 24,
            "Fonction bénéficiaire": 24,
            "N° externe / DEM Jira": 22,
            "Statut JIRA": 18,
            "Localisation (dernier niveau)": 32,
            "Relance(s)": 18,
            "Date d'émission": 21,
            "Date de dernière modification": 24,
            "Date de planification": 21,
            "Date de résolution maximum de l'incident": 26,
            "Date de fin réelle": 21,
        }

        for column_index, header in enumerate(headers):
            ws.set_column(column_index, column_index, widths.get(header, 18))

        ws.set_row(0, 26)

        # Ajoute une couleur alternée légère pour faciliter la lecture, sans
        # modifier la structure des données.
        if tickets and headers:
            last_row = len(tickets)
            last_col = len(headers) - 1
            ws.conditional_format(
                1,
                0,
                last_row,
                last_col,
                {
                    "type": "formula",
                    "criteria": "=MOD(ROW(),2)=0",
                    "format": workbook.add_format({"bg_color": "#F7FAFC"}),
                },
            )

        # ---------------------------------------------------------------------
        # 7.3) Onglet "Dashboard"
        # ---------------------------------------------------------------------
        dash.merge_range(
            "A1:H1",
            "Dashboard EasyVista — tickets uniques",
            fmt_title,
        )
        dash.set_row(0, 28)

        # KPI principal : nombre de tickets uniques.
        dash.write("A3", "Indicateur", fmt_section)
        dash.write("B3", "Valeur", fmt_section)
        dash.write("A4", "Tickets uniques")
        dash.write("B4", len(tickets), fmt_kpi)

        status, priority, types, groups, months = build_dashboard_data(tickets)

        def write_summary(row, col, title, counter, sort_mode="count", limit=None):
            """Ecrit un petit tableau de synthèse dans le Dashboard.

            row / col sont des index XlsxWriter commençant à 0.
            La fonction retourne les bornes de données utilisées par le graphique.
            """

            items = list(counter.items())

            if sort_mode == "key":
                items.sort(key=lambda item: str(item[0]))
            elif sort_mode == "natural":
                items.sort(key=lambda item: natural_key(item[0]))
            else:
                # Par défaut : catégories les plus fréquentes en premier.
                items.sort(key=lambda item: (-item[1], str(item[0])))

            if limit:
                items = items[:limit]

            dash.write(row, col, title, fmt_section)
            dash.write(row, col + 1, "Tickets", fmt_section)

            for offset, (label, count) in enumerate(items, start=1):
                # Pour les mois, transforme 2026-08 en 08/2026 pour l'affichage.
                display_label = str(label)
                if title == "Mois" and re.fullmatch(r"\d{4}-\d{2}", display_label):
                    year, month = display_label.split("-")
                    display_label = f"{month}/{year}"

                dash.write(row + offset, col, display_label)
                dash.write_number(row + offset, col + 1, int(count))

            # Première et dernière ligne de DONNEES, hors en-tête.
            return row + 1, row + len(items)

        # Tableaux servant à la fois de synthèse lisible et de source aux charts.
        s1, e1 = write_summary(6, 0, "Statut", status)
        s2, e2 = write_summary(6, 3, "Priorité", priority, sort_mode="natural")
        s3, e3 = write_summary(6, 6, "Type", types)
        s4, e4 = write_summary(20, 0, "Mois", months, sort_mode="key")
        s5, e5 = write_summary(20, 3, "Top groupes", groups, limit=10)

        def add_chart(kind, title, cat_col, val_col, start, end, cell,
                      legend=False, data_labels=False):
            """Ajoute un graphique à partir d'un tableau du Dashboard."""

            # Si le tableau source est vide, aucun graphique n'est créé.
            if end < start:
                return

            chart = workbook.add_chart({"type": kind})

            series = {
                "name": title,
                "categories": ["Dashboard", start, cat_col, end, cat_col],
                "values": ["Dashboard", start, val_col, end, val_col],
            }

            if data_labels:
                series["data_labels"] = {"value": True}

            chart.add_series(series)
            chart.set_title({"name": title})
            chart.set_style(10)

            if legend:
                chart.set_legend({"position": "bottom"})
            else:
                chart.set_legend({"none": True})

            # Taille uniforme pour garder un Dashboard propre.
            dash.insert_chart(
                cell,
                chart,
                {"x_scale": 1.18, "y_scale": 1.12},
            )

        # Graphiques basés uniquement sur les TICKETS UNIQUES.
        add_chart(
            "column",
            "Tickets uniques par statut",
            0, 1, s1, e1, "J2",
            data_labels=True,
        )

        add_chart(
            "column",
            "Tickets uniques par priorité",
            3, 4, s2, e2, "J18",
            data_labels=True,
        )

        # Le camembert est construit séparément afin d'afficher catégorie + %.
        if e3 >= s3:
            chart_type = workbook.add_chart({"type": "pie"})
            chart_type.add_series({
                "name": "Incidents / Demandes de service",
                "categories": ["Dashboard", s3, 6, e3, 6],
                "values": ["Dashboard", s3, 7, e3, 7],
                "data_labels": {
                    "category": True,
                    "percentage": True,
                    "leader_lines": True,
                },
            })
            chart_type.set_title({"name": "Incidents / Demandes de service"})
            chart_type.set_legend({"position": "bottom"})
            chart_type.set_style(10)
            dash.insert_chart(
                "R2",
                chart_type,
                {"x_scale": 1.18, "y_scale": 1.12},
            )

        add_chart(
            "line",
            "Tickets uniques par mois",
            0, 1, s4, e4, "R18",
            data_labels=True,
        )

        add_chart(
            "bar",
            "Top groupes — tickets uniques",
            3, 4, s5, e5, "J34",
            data_labels=True,
        )

        dash.set_column("A:H", 20)
        dash.freeze_panes(1, 0)

        # Fermer le workbook écrit réellement le fichier .xlsx sur disque.
        workbook.close()

    except FileCreateError as exc:
        # Message fréquent sous Windows lorsque le fichier cible est déjà
        # ouvert dans Excel et donc verrouillé.
        raise RuntimeError(
            f"Impossible de créer '{output}'. "
            "S'il est déjà ouvert dans Excel, ferme-le puis relance le script."
        ) from exc


# =============================================================================
# 8) PROGRAMME PRINCIPAL / OPTIONS EN LIGNE DE COMMANDE
# =============================================================================

def main():
    """Point d'entrée du script."""

    parser = argparse.ArgumentParser(
        description=(
            "EasyVista : transforme 'Mes actions' en une liste de tickets "
            "uniques avec Dashboard Excel."
        )
    )

    parser.add_argument(
        "fichier",
        type=Path,
        help="Export EasyVista au format .csv ou .xlsx",
    )

    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        help=(
            "Fichier .xlsx de sortie. Par défaut : "
            "<nom_source>_tickets_uniques_dashboard.xlsx"
        ),
    )

    parser.add_argument(
        "--sheet",
        help="Nom de la feuille à lire si le fichier source est un XLSX.",
    )

    parser.add_argument(
        "-v",
        "--verbose",
        action="store_true",
        help="Affiche des informations techniques sur le traitement.",
    )

    args = parser.parse_args()
    source = args.fichier

    # -------------------------------------------------------------------------
    # Vérifications avant traitement
    # -------------------------------------------------------------------------
    if not source.exists():
        print(f"Erreur : fichier introuvable : {source}", file=sys.stderr)
        return 1

    if not source.is_file():
        print(f"Erreur : ce chemin n'est pas un fichier : {source}", file=sys.stderr)
        return 1

    output = args.output or source.with_name(
        f"{source.stem}_tickets_uniques_dashboard.xlsx"
    )

    # On force une sortie XLSX : xlsxwriter crée uniquement ce format.
    if output.suffix.lower() != ".xlsx":
        output = output.with_suffix(".xlsx")

    # Evite d'écraser le fichier source si l'utilisateur passe le même nom.
    try:
        same_file = source.resolve() == output.resolve()
    except OSError:
        same_file = False

    if same_file:
        print(
            "Erreur : le fichier de sortie ne peut pas être le fichier source.",
            file=sys.stderr,
        )
        return 1

    # -------------------------------------------------------------------------
    # Traitement complet
    # -------------------------------------------------------------------------
    try:
        # ETAPE 1 : lecture CSV/XLSX.
        headers, rows = read_export(
            source,
            sheet_name=args.sheet,
            verbose=args.verbose,
        )

        # ETAPE 2 : suppression de la ligne technique et des doublons de ticket.
        actions, tickets = deduplicate(
            headers,
            rows,
            verbose=args.verbose,
        )

        if not tickets:
            raise ValueError("Aucun ticket valide n'a été trouvé dans l'export.")

        # ETAPE 3 : création du fichier Excel final.
        write_output(headers, tickets, output)

        # ---------------------------------------------------------------------
        # Résumé affiché dans la console
        # ---------------------------------------------------------------------
        print()
        print("Traitement terminé avec succès.")
        print(f"Lignes/actions analysées : {len(actions)}")
        print(f"Tickets uniques          : {len(tickets)}")
        print(f"Doublons retirés         : {len(actions) - len(tickets)}")
        print(f"Fichier créé             : {output.resolve()}")

        return 0

    except Exception as exc:
        print(f"Erreur : {exc}", file=sys.stderr)
        return 2


# =============================================================================
# 9) LANCEMENT DU SCRIPT
# =============================================================================
# Ce bloc permet d'exécuter main() uniquement quand le fichier est lancé
# directement avec "python easyvista_tickets.py ...".
# Si le fichier était importé comme module dans un autre script, main() ne serait
# pas exécuté automatiquement.

if __name__ == "__main__":
    raise SystemExit(main())
